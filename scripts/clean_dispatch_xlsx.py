#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
清洗并美化「一键派单」导出的 Excel 派单表。

功能：
1. 清洗「产品描述」列：去除所有 HTML 标签（<p>、<br/>、<img src="..."/> 等），
   只保留纯文字内容。
2. 美化格式：
   - 「产品描述」列宽按文字长度自适应；
   - 该列单元格开启自动换行；
   - 数据行取消固定行高，交由 WPS/Excel 按换行后内容自动撑开（避免文字被截断）。
3. 另存为全新文件（不修改原文件）：<原文件名>_清洗后.xlsx
   文件中的图片（如「做单文件码」二维码）会被完整保留。

依赖安装（pillow 必需，缺少时 openpyxl 保存会丢失二维码图片）：
    pip3 install openpyxl pillow

用法：
    python3 clean_dispatch_xlsx.py [派单表_2026-10-24.xlsx]
    （不传参数时默认处理当前目录下的 派单表_2026-10-24.xlsx）
"""

import html
import re
import sys
from pathlib import Path

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

try:
    import openpyxl
    import PIL  # noqa: F401  没有 Pillow 时 openpyxl 读存会丢图
except ImportError as e:
    sys.exit(f"缺少依赖库：{e.name}，请先执行  pip3 install openpyxl pillow")

DEFAULT_INPUT = "派单表_2026-10-24.xlsx"
DESC_HEADER = "产品描述"
MIN_COL_WIDTH = 20   # 描述列最小列宽
MAX_COL_WIDTH = 80   # 描述列最大列宽（超出部分交给自动换行）


# ---------------------------------------------------------------- HTML 清洗 --
def clean_html_text(value) -> object:
    """去除 HTML 标签与图片路径代码，仅保留纯文字。非字符串原样返回。"""
    if not isinstance(value, str) or not value:
        return value
    text = re.sub(r"<img[^>]*/?>", "", value, flags=re.I)      # 去图片代码
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)        # 换行标签转真实换行
    text = re.sub(r"<[^>]+>", "", text)                        # 去其余所有标签
    text = html.unescape(text)                                 # 还原 &nbsp; 等实体
    text = text.replace("\xa0", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line)           # 去空行、去首尾空白


def display_width(text: str) -> int:
    """估算文字显示宽度：中文等全角字符按 2 个单位计。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in text)


# ---------------------------------------------------------------- 主流程 --
def process(input_path: Path) -> Path:
    wb = openpyxl.load_workbook(input_path)

    for ws in wb.worksheets:
        # 1. 定位所有「产品描述」表头，确定列号与表头所在行
        header_rows, desc_col = set(), None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == DESC_HEADER:
                    header_rows.add(cell.row)
                    desc_col = cell.column
        if desc_col is None:
            print(f"[跳过] 工作表「{ws.title}」未找到「{DESC_HEADER}」列")
            continue

        # 2. 清洗描述列数据（跳过表头行）
        max_width, cleaned_rows, html_rows = 0, [], 0
        for r in range(1, ws.max_row + 1):
            if r in header_rows:
                continue
            cell = ws.cell(row=r, column=desc_col)
            if not isinstance(cell.value, str) or not cell.value:
                continue
            if "<" in cell.value:
                html_rows += 1
            cell.value = clean_html_text(cell.value)
            max_width = max(max_width, max(display_width(line) for line in cell.value.splitlines()))
            # 3. 开启自动换行（保留原有对齐方式）
            old = cell.alignment
            cell.alignment = Alignment(
                horizontal=old.horizontal, vertical=old.vertical or "center",
                wrap_text=True, shrink_to_fit=False,
            )
            cleaned_rows.append(r)

        # 4. 描述列宽自适应（限制上下限，超宽内容由自动换行承接）
        if max_width:
            ws.column_dimensions[get_column_letter(desc_col)].width = \
                min(max(max_width + 2, MIN_COL_WIDTH), MAX_COL_WIDTH)

        # 5. 数据行取消固定行高，让 WPS/Excel 自动按内容撑开
        for r in cleaned_rows:
            if r in ws.row_dimensions:
                ws.row_dimensions[r].height = None

        print(f"[完成] 工作表「{ws.title}」：清洗 {len(cleaned_rows)} 行"
              f"（其中含 HTML 标签 {html_rows} 行），"
              f"描述列宽设为 {ws.column_dimensions[get_column_letter(desc_col)].width:.0f}，"
              f"保留图片 {len(ws._images)} 张")

    output_path = input_path.with_name(f"{input_path.stem}_清洗后.xlsx")
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(DEFAULT_INPUT)
    if not src.exists():
        sys.exit(f"未找到文件：{src.resolve()}")
    out = process(src)
    print(f"已保存：{out}")
